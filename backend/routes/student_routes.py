from flask import Blueprint, request, jsonify
from flask_jwt_extended import get_jwt_identity, get_jwt
from middleware.auth_middleware import teacher_required, student_required, validate_standard
from models.student_model import (
    get_students_by_standard, get_student_by_id,
    create_student, update_student, delete_student
)
from pymongo.errors import DuplicateKeyError

student_bp = Blueprint('students', __name__)

@student_bp.route('', methods=['GET'])
@teacher_required
def list_students():
    std_param = request.args.get('standard')
    if not std_param:
        return jsonify({'success': False, 'error': 'standard query parameter is required.'}), 400
    std, err = validate_standard(std_param)
    if err:
        return err
    students = get_students_by_standard(std)
    return jsonify({'success': True, 'standard': std, 'students': students, 'count': len(students)})

@student_bp.route('/<student_id>', methods=['GET'])
@teacher_required
def get_student(student_id):
    student = get_student_by_id(student_id)
    if not student:
        return jsonify({'success': False, 'error': 'Student not found.'}), 404
    return jsonify({'success': True, 'student': student})

@student_bp.route('', methods=['POST'])
@teacher_required
def add_student():
    data = request.get_json(silent=True) or {}
    required = ['name', 'roll_no', 'standard', 'password']
    missing = [f for f in required if not data.get(f)]
    if missing:
        return jsonify({'success': False, 'error': f"Missing fields: {', '.join(missing)}"}), 400

    std, err = validate_standard(data['standard'])
    if err:
        return err
    data['standard'] = std

    try:
        student = create_student(data)
        return jsonify({'success': True, 'student': student}), 201
    except DuplicateKeyError as e:
        key = 'roll number' if 'roll_no' in str(e) else 'email'
        return jsonify({'success': False, 'error': f'A student with that {key} already exists.'}), 409

@student_bp.route('/<student_id>', methods=['PUT'])
@teacher_required
def edit_student(student_id):
    data = request.get_json(silent=True) or {}
    student = get_student_by_id(student_id)
    if not student:
        return jsonify({'success': False, 'error': 'Student not found.'}), 404
    try:
        updated = update_student(student_id, data)
        return jsonify({'success': True, 'student': updated})
    except DuplicateKeyError as e:
        key = 'roll number' if 'roll_no' in str(e) else 'email'
        return jsonify({'success': False, 'error': f'A student with that {key} already exists.'}), 409

@student_bp.route('/<student_id>', methods=['DELETE'])
@teacher_required
def remove_student(student_id):
    if not get_student_by_id(student_id):
        return jsonify({'success': False, 'error': 'Student not found.'}), 404
    delete_student(student_id)
    return jsonify({'success': True, 'message': 'Student deleted successfully.'})

@student_bp.route('/me', methods=['GET'])
@student_required
def my_profile():
    student_id = get_jwt_identity()
    student = get_student_by_id(student_id)
    if not student:
        return jsonify({'success': False, 'error': 'Student not found.'}), 404
    return jsonify({'success': True, 'student': student})

@student_bp.route('/me/photo', methods=['POST', 'PUT'])
@student_required
def update_my_photo():
    student_id = get_jwt_identity()
    data = request.get_json(silent=True) or {}
    photo = data.get('photo') or data.get('photo_url')
    if not photo:
        return jsonify({'success': False, 'error': 'Photo image or URL is required.'}), 400
    updated = update_student(student_id, {'photo': photo})
    return jsonify({'success': True, 'message': 'Profile photo updated successfully.', 'student': updated})


# ─── Excel / CSV Template, Import & Export ───────────────────────────────────

import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import send_file

@student_bp.route('/template', methods=['GET'])
@teacher_required
def download_student_template():
    """Generates and returns a formatted Excel template for bulk student import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Import Template"

    # Define styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    example_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid") # Light Green
    example_font = Font(name="Arial", size=10, italic=True, color="166534")
    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = ["Roll No *", "Full Name *", "Standard (1-12) *", "Mobile Number (10 Digits)", "Email Address", "Password (Default: student123)"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    # Sample demo rows for user guidance
    samples = [
        [501, "Aarav Patel", 5, "9876505001", "aarav.p@school.com", "student123"],
        [502, "Diya Sharma", 5, "9876505002", "diya.s@school.com", "student123"],
        [503, "Rohan Kumar", 5, "9876505003", "rohan.k@school.com", "student123"],
    ]

    for row_data in samples:
        ws.append(row_data)
        row_idx = ws.max_row
        for col_num in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=col_num)
            c.fill = example_fill
            c.font = example_font
            c.border = border_thin
            if col_num in [1, 3, 4]:
                c.alignment = Alignment(horizontal="center")

    # Column widths
    col_widths = [15, 28, 20, 26, 30, 32]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="student_import_template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@student_bp.route('/import', methods=['POST'])
@teacher_required
def import_students():
    """Bulk import students from an Excel (.xlsx) or CSV file or JSON body."""
    records = []

    if 'file' in request.files:
        file = request.files['file']
        filename = (file.filename or '').lower()

        if filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode('utf-8-sig', errors='ignore'))
            reader = csv.DictReader(stream)
            for row in reader:
                records.append(row)
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return jsonify({'success': False, 'error': 'Excel file is empty.'}), 400

            header_row = [str(c).strip().lower() if c is not None else '' for c in rows[0]]
            for r in rows[1:]:
                if not any(r):
                    continue
                row_dict = {}
                for idx, val in enumerate(r):
                    if idx < len(header_row):
                        key = header_row[idx]
                        row_dict[key] = str(val).strip() if val is not None else ''
                records.append(row_dict)
        else:
            return jsonify({'success': False, 'error': 'Unsupported file type. Please upload a .xlsx or .csv file.'}), 400
    else:
        # Check JSON payload
        data = request.get_json(silent=True) or {}
        records = data.get('students', data if isinstance(data, list) else [])

    if not records:
        return jsonify({'success': False, 'error': 'No student records found in uploaded file or payload.'}), 400

    imported_count = 0
    skipped_count = 0
    errors = []

    for idx, row in enumerate(records, start=2):
        # Flexible header mapping
        roll_val = None
        for k in ['roll no *', 'roll no', 'roll_no', 'roll', 'rollno']:
            if k in row and row[k]:
                roll_val = row[k]
                break

        name_val = None
        for k in ['full name *', 'full name', 'name', 'student name', 'student_name']:
            if k in row and row[k]:
                name_val = row[k]
                break

        std_val = None
        for k in ['standard (1-12) *', 'standard', 'std', 'class']:
            if k in row and row[k]:
                std_val = row[k]
                break

        mobile_val = None
        for k in ['mobile number (10 digits)', 'mobile', 'mobile_no', 'phone', 'contact']:
            if k in row and row[k]:
                mobile_val = str(row[k]).replace('.0', '').strip()
                break

        email_val = None
        for k in ['email address', 'email', 'e-mail']:
            if k in row and row[k]:
                email_val = str(row[k]).strip()
                break

        pwd_val = None
        for k in ['password (default: student123)', 'password', 'pwd']:
            if k in row and row[k]:
                pwd_val = str(row[k]).strip()
                break

        # Validate required fields
        if not roll_val or not name_val or not std_val:
            skipped_count += 1
            errors.append(f"Row {idx}: Missing required fields (Roll No, Name, or Standard).")
            continue

        try:
            roll_int = int(float(str(roll_val).strip()))
        except ValueError:
            skipped_count += 1
            errors.append(f"Row {idx}: Invalid Roll No '{roll_val}'. Must be a number.")
            continue

        try:
            std_int = int(float(str(std_val).strip()))
            if not (1 <= std_int <= 12):
                raise ValueError()
        except ValueError:
            skipped_count += 1
            errors.append(f"Row {idx}: Standard must be between 1 and 12 (got '{std_val}').")
            continue

        student_payload = {
            'name': str(name_val).strip(),
            'roll_no': roll_int,
            'standard': std_int,
            'password': pwd_val or 'student123',
            'mobile': mobile_val or None,
            'email': email_val or None
        }

        try:
            create_student(student_payload)
            imported_count += 1
        except DuplicateKeyError as e:
            skipped_count += 1
            err_msg = "Roll number" if "roll_no" in str(e) else "Mobile or Email"
            errors.append(f"Row {idx} ({student_payload['name']}): Duplicate {err_msg} already exists.")
        except Exception as e:
            skipped_count += 1
            errors.append(f"Row {idx} ({student_payload['name']}): {str(e)}")

    return jsonify({
        'success': True,
        'message': f"Import complete: {imported_count} student(s) added, {skipped_count} skipped.",
        'imported_count': imported_count,
        'skipped_count': skipped_count,
        'errors': errors[:20] # Top 20 errors to prevent payload explosion
    }), 200


@student_bp.route('/export', methods=['GET'])
@teacher_required
def export_students():
    """Exports students directory to a styled Excel (.xlsx) file."""
    std_param = request.args.get('standard')
    if std_param:
        std, err = validate_standard(std_param)
        if err:
            return err
        students = get_students_by_standard(std)
        sheet_title = f"Standard {std} Students"
        file_name = f"students_standard_{std}.xlsx"
    else:
        # Export all students across standards
        from database.db import get_db
        db = get_db()
        docs = db.students.find({}).sort([('standard', 1), ('roll_no', 1)])
        from models.student_model import serialize
        students = [serialize(d) for d in docs]
        sheet_title = "All Students Directory"
        file_name = "all_students_directory.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Title Banner
    ws.merge_cells('A1:F1')
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"EDUMANAGE PRO - {sheet_title.upper()}"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Headers
    headers = ["Roll No", "Full Name", "Standard", "Mobile Number", "Email Address", "Registered Date"]
    ws.append(headers)

    header_font = Font(name="Arial", size=10, bold=True, color="1E3A8A")
    header_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.row_dimensions[2].height = 24
    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=2, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_thin

    # Data Rows
    row_font = Font(name="Arial", size=10)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, s in enumerate(students, start=3):
        created_str = str(s.get('created_at', ''))[:10] if s.get('created_at') else 'N/A'
        row_vals = [
            s.get('roll_no', ''),
            s.get('name', ''),
            f"Std {s.get('standard', '')}",
            s.get('mobile') or '—',
            s.get('email') or '—',
            created_str
        ]
        ws.append(row_vals)
        ws.row_dimensions[idx].height = 20
        use_fill = (idx % 2 == 0)

        for col_num in range(1, len(row_vals) + 1):
            c = ws.cell(row=idx, column=col_num)
            c.font = row_font
            c.border = border_thin
            if use_fill:
                c.fill = alt_fill
            if col_num in [1, 3, 4, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    # Column Widths
    widths = [12, 28, 14, 20, 28, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=file_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
